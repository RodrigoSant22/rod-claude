"""Exportação para CSV e Excel."""

import csv
import io
from datetime import date
from decimal import Decimal

from openpyxl import load_workbook

from app import exportacao


def ler_csv(conteudo: bytes) -> list[list[str]]:
    """Relê o CSV gerado, descartando o BOM, como o Excel faria."""
    texto = conteudo.decode("utf-8-sig")
    return list(csv.reader(io.StringIO(texto), delimiter=";"))


def ler_xlsx(conteudo: bytes):
    """Relê a planilha gerada, para inspecionar células e formatos."""
    return load_workbook(io.BytesIO(conteudo))


# --------------------------------------------------------------------------
# CSV
# --------------------------------------------------------------------------


def test_csv_tem_bom_para_o_excel(app, lancamentos):
    """Sem BOM, o Excel abre o arquivo com os acentos quebrados."""
    conteudo = exportacao.gerar_csv(lancamentos)

    assert conteudo.startswith(b"\xef\xbb\xbf")


def test_csv_usa_ponto_e_virgula(app, lancamentos):
    """Com vírgula, o Excel em português joga tudo numa coluna só."""
    linhas = ler_csv(exportacao.gerar_csv(lancamentos))

    assert linhas[0] == ["Data", "Tipo", "Categoria", "Descrição", "Valor", "Observação"]


def test_csv_traz_todos_os_lancamentos(app, lancamentos):
    linhas = ler_csv(exportacao.gerar_csv(lancamentos))

    assert len(linhas) == len(lancamentos) + 1  # + cabeçalho


def test_csv_usa_virgula_decimal(app, lancamentos):
    linhas = ler_csv(exportacao.gerar_csv(lancamentos))
    valores = [linha[4] for linha in linhas[1:]]

    assert all("," in valor for valor in valores)
    assert all("." not in valor for valor in valores)


def test_csv_traz_despesa_negativa(app, lancamentos):
    """Uma coluna com sinal permite SOMA() direto na planilha."""
    linhas = ler_csv(exportacao.gerar_csv(lancamentos))
    por_descricao = {linha[3]: linha for linha in linhas[1:]}

    assert por_descricao["Salário de março"][4] == "5000,00"
    assert por_descricao["Cinema"][4] == "-64,90"


def test_csv_data_em_formato_brasileiro(app, lancamentos):
    linhas = ler_csv(exportacao.gerar_csv(lancamentos))
    por_descricao = {linha[3]: linha for linha in linhas[1:]}

    assert por_descricao["Salário de março"][0] == "05/03/2026"


def test_csv_sem_lancamentos_traz_so_o_cabecalho(app):
    linhas = ler_csv(exportacao.gerar_csv([]))

    assert len(linhas) == 1


def test_csv_preserva_acentos(app, lancamentos):
    linhas = ler_csv(exportacao.gerar_csv(lancamentos))

    assert any("Salário de março" in linha for linha in linhas)


# --------------------------------------------------------------------------
# XLSX
# --------------------------------------------------------------------------


def test_xlsx_tem_as_duas_abas(app, lancamentos):
    livro = ler_xlsx(exportacao.gerar_xlsx(lancamentos))

    assert livro.sheetnames == ["Lançamentos", "Resumo"]


def test_xlsx_grava_valor_como_numero(app, lancamentos):
    """Como texto, a planilha não somaria nem filtraria por faixa."""
    aba = ler_xlsx(exportacao.gerar_xlsx(lancamentos))["Lançamentos"]

    for linha in range(2, aba.max_row + 1):
        valor = aba.cell(row=linha, column=5).value
        assert isinstance(valor, int | float | Decimal), f"linha {linha}: {type(valor)}"


def test_xlsx_grava_data_como_data(app, lancamentos):
    aba = ler_xlsx(exportacao.gerar_xlsx(lancamentos))["Lançamentos"]

    valor = aba.cell(row=2, column=1).value
    assert hasattr(valor, "year")


def test_xlsx_aplica_formato_de_moeda(app, lancamentos):
    aba = ler_xlsx(exportacao.gerar_xlsx(lancamentos))["Lançamentos"]

    assert "R$" in aba.cell(row=2, column=5).number_format


def test_xlsx_congela_o_cabecalho_e_liga_o_filtro(app, lancamentos):
    aba = ler_xlsx(exportacao.gerar_xlsx(lancamentos))["Lançamentos"]

    assert aba.freeze_panes == "A2"
    assert aba.auto_filter.ref is not None


def test_xlsx_despesa_sai_negativa(app, lancamentos):
    aba = ler_xlsx(exportacao.gerar_xlsx(lancamentos))["Lançamentos"]

    valores = {
        aba.cell(row=n, column=4).value: aba.cell(row=n, column=5).value
        for n in range(2, aba.max_row + 1)
    }

    assert valores["Salário de março"] == 5000
    assert valores["Cinema"] == -64.9


def test_xlsx_resumo_bate_com_os_dados(app, lancamentos):
    aba = ler_xlsx(exportacao.gerar_xlsx(lancamentos))["Resumo"]
    indicadores = {
        aba.cell(row=n, column=1).value: aba.cell(row=n, column=2).value
        for n in range(2, 6)
    }

    assert indicadores["Receitas"] == 5000
    assert indicadores["Despesas"] == 3664.9
    assert round(indicadores["Saldo"], 2) == 1335.1
    assert indicadores["Lançamentos"] == 4


def test_xlsx_resumo_agrupa_por_categoria(app, lancamentos):
    aba = ler_xlsx(exportacao.gerar_xlsx(lancamentos))["Resumo"]
    linhas = [
        (aba.cell(row=n, column=1).value, aba.cell(row=n, column=3).value)
        for n in range(2, aba.max_row + 1)
    ]

    assert ("Moradia", 3600) in linhas


def test_xlsx_sem_lancamentos_nao_quebra(app):
    livro = ler_xlsx(exportacao.gerar_xlsx([]))

    assert livro["Lançamentos"].max_row == 1
    assert livro["Resumo"]["B2"].value == 0


# --------------------------------------------------------------------------
# Nome do arquivo
# --------------------------------------------------------------------------


def test_nome_com_periodo_completo():
    nome = exportacao.nome_arquivo("xlsx", date(2026, 3, 1), date(2026, 3, 31))

    assert nome == "lancamentos-2026-03-01-a-2026-03-31.xlsx"


def test_nome_sem_periodo_usa_hoje():
    nome = exportacao.nome_arquivo("csv", None, None)

    assert nome == f"lancamentos-{date.today().isoformat()}.csv"


def test_nome_com_apenas_um_limite():
    assert exportacao.nome_arquivo("csv", date(2026, 3, 1), None).endswith(
        "desde-2026-03-01.csv"
    )
    assert exportacao.nome_arquivo("csv", None, date(2026, 3, 31)).endswith(
        "ate-2026-03-31.csv"
    )

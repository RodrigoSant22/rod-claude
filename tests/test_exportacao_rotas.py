"""Rotas de exportação: autenticação, filtros e isolamento entre contas."""

import csv
import io


def baixar(cliente, caminho: str):
    """Baixa um arquivo e falha cedo se a rota não devolveu 200."""
    resposta = cliente.get(caminho)
    assert resposta.status_code == 200, caminho
    return resposta


def linhas_csv(resposta) -> list[list[str]]:
    """Linhas de dados do CSV baixado, sem o cabeçalho."""
    texto = resposta.get_data().decode("utf-8-sig")
    return list(csv.reader(io.StringIO(texto), delimiter=";"))[1:]


def test_exportar_exige_login(client):
    for formato in ("csv", "xlsx"):
        resp = client.get(f"/lancamentos/exportar.{formato}")
        assert resp.status_code == 302
        assert "/login" in resp.headers["Location"]


def test_csv_tem_cabecalhos_de_download(logado, lancamentos):
    resp = baixar(logado, "/lancamentos/exportar.csv")

    assert resp.mimetype == "text/csv"
    assert "attachment" in resp.headers["Content-Disposition"]
    assert ".csv" in resp.headers["Content-Disposition"]


def test_xlsx_tem_cabecalhos_de_download(logado, lancamentos):
    resp = baixar(logado, "/lancamentos/exportar.xlsx")

    assert "spreadsheetml" in resp.mimetype
    assert ".xlsx" in resp.headers["Content-Disposition"]
    # Todo xlsx é um zip: começa com "PK".
    assert resp.get_data()[:2] == b"PK"


def test_formato_desconhecido_da_404(logado, lancamentos):
    assert logado.get("/lancamentos/exportar.pdf").status_code == 404


def test_exportacao_respeita_o_filtro_de_texto(logado, lancamentos):
    resp = baixar(logado, "/lancamentos/exportar.csv?texto=cinema")
    dados = linhas_csv(resp)

    assert len(dados) == 1
    assert dados[0][3] == "Cinema"


def test_exportacao_respeita_o_filtro_de_periodo(logado, lancamentos):
    resp = baixar(logado, "/lancamentos/exportar.csv?inicio=2026-04-01&fim=2026-04-30")
    dados = linhas_csv(resp)

    assert len(dados) == 1
    assert dados[0][0] == "10/04/2026"


def test_exportacao_respeita_o_filtro_de_tipo(logado, lancamentos):
    resp = baixar(logado, "/lancamentos/exportar.csv?tipo=receita")
    dados = linhas_csv(resp)

    assert len(dados) == 1
    assert dados[0][1] == "Receita"


def test_nome_do_arquivo_reflete_o_periodo(logado, lancamentos):
    resp = baixar(logado, "/lancamentos/exportar.csv?inicio=2026-03-01&fim=2026-03-31")

    assert "lancamentos-2026-03-01-a-2026-03-31.csv" in resp.headers["Content-Disposition"]


def test_exportacao_nao_vaza_dados_de_outra_conta(logado, lancamentos, dados_do_outro):
    resp = baixar(logado, "/lancamentos/exportar.csv")
    conteudo = resp.get_data().decode("utf-8-sig")

    assert "Salário de março" in conteudo
    assert "Passagem aérea secreta" not in conteudo


def test_xlsx_nao_vaza_dados_de_outra_conta(logado, lancamentos, dados_do_outro):
    from openpyxl import load_workbook

    resp = baixar(logado, "/lancamentos/exportar.xlsx")
    aba = load_workbook(io.BytesIO(resp.get_data()))["Lançamentos"]

    descricoes = {aba.cell(row=n, column=4).value for n in range(2, aba.max_row + 1)}
    assert "Salário de março" in descricoes
    assert "Passagem aérea secreta" not in descricoes


def test_listagem_mostra_os_links_com_os_filtros(logado, lancamentos):
    corpo = logado.get("/lancamentos/?texto=cinema").get_data(as_text=True)

    assert "exportar.xlsx" in corpo
    assert "texto=cinema" in corpo


def test_links_acompanham_o_fragmento_do_htmx(logado, lancamentos):
    """Os links ficam dentro do pedaço que o HTMX troca, senão guardariam
    os filtros antigos depois de filtrar."""
    corpo = logado.get(
        "/lancamentos/?tipo=receita", headers={"HX-Request": "true"}
    ).get_data(as_text=True)

    assert "exportar.csv" in corpo
    assert "tipo=receita" in corpo


def test_sem_resultados_nao_mostra_o_convite(logado, lancamentos):
    corpo = logado.get("/lancamentos/?texto=nada-disso").get_data(as_text=True)

    assert "Exportar" not in corpo

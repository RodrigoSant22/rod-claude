"""Geração dos arquivos de exportação.

Duas escolhas guiam o módulo:

**O valor sai com sinal** — negativo para despesa. Uma coluna numérica só
permite `SOMA()` direto e funciona em tabela dinâmica; duas colunas
separadas obrigariam a subtrair uma da outra a cada análise.

**CSV pensado para o Excel em português** — separador `;` e UTF-8 com BOM.
Sem isso, abrir o arquivo com duplo clique no Brasil joga tudo numa coluna
só e estraga os acentos.
"""

import csv
import io
from datetime import date
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.models import Lancamento, TipoLancamento

CABECALHO = ["Data", "Tipo", "Categoria", "Descrição", "Valor", "Observação"]

# Larguras em caracteres, na ordem do cabeçalho.
LARGURAS = [12, 10, 22, 40, 14, 40]

FORMATO_MOEDA = 'R$ #,##0.00;[Red]-R$ #,##0.00'
FORMATO_DATA = "DD/MM/YYYY"


def _linhas(lancamentos: list[Lancamento]) -> list[tuple]:
    return [
        (
            item.data,
            item.tipo.rotulo,
            item.categoria.nome,
            item.descricao,
            item.valor_com_sinal,
            item.observacao or "",
        )
        for item in lancamentos
    ]


def nome_arquivo(extensao: str, inicio: date | None, fim: date | None) -> str:
    """lancamentos-2026-03-01-a-2026-03-31.xlsx, ou com a data de hoje."""
    if inicio and fim:
        miolo = f"{inicio.isoformat()}-a-{fim.isoformat()}"
    elif inicio:
        miolo = f"desde-{inicio.isoformat()}"
    elif fim:
        miolo = f"ate-{fim.isoformat()}"
    else:
        miolo = date.today().isoformat()
    return f"lancamentos-{miolo}.{extensao}"


# --------------------------------------------------------------------------
# CSV
# --------------------------------------------------------------------------


def gerar_csv(lancamentos: list[Lancamento]) -> bytes:
    buffer = io.StringIO()
    escritor = csv.writer(buffer, delimiter=";", lineterminator="\r\n")

    escritor.writerow(CABECALHO)
    for data, tipo, categoria, descricao, valor, observacao in _linhas(lancamentos):
        escritor.writerow(
            [
                data.strftime("%d/%m/%Y"),
                tipo,
                categoria,
                descricao,
                # Vírgula decimal: é o que o Excel em português entende como
                # número. Com ponto, ele trataria o valor como texto.
                f"{valor:.2f}".replace(".", ","),
                observacao,
            ]
        )

    # utf-8-sig acrescenta o BOM que o Excel usa para detectar a codificação.
    return buffer.getvalue().encode("utf-8-sig")


# --------------------------------------------------------------------------
# XLSX
# --------------------------------------------------------------------------

_CABECALHO_FUNDO = PatternFill("solid", fgColor="1F2937")
_CABECALHO_FONTE = Font(color="FFFFFF", bold=True)


def _escrever_cabecalho(planilha, colunas: list[str]) -> None:
    planilha.append(colunas)
    for coluna in range(1, len(colunas) + 1):
        celula = planilha.cell(row=1, column=coluna)
        celula.fill = _CABECALHO_FUNDO
        celula.font = _CABECALHO_FONTE
        celula.alignment = Alignment(vertical="center")
    planilha.row_dimensions[1].height = 20


def _aba_lancamentos(planilha, lancamentos: list[Lancamento]) -> None:
    planilha.title = "Lançamentos"
    _escrever_cabecalho(planilha, CABECALHO)

    for linha in _linhas(lancamentos):
        planilha.append(linha)

    for indice, largura in enumerate(LARGURAS, start=1):
        planilha.column_dimensions[get_column_letter(indice)].width = largura

    ultima = planilha.max_row
    for linha in range(2, ultima + 1):
        planilha.cell(row=linha, column=1).number_format = FORMATO_DATA
        planilha.cell(row=linha, column=5).number_format = FORMATO_MOEDA

    # Congela o cabeçalho e liga o filtro, para a planilha já chegar usável.
    planilha.freeze_panes = "A2"
    if ultima > 1:
        planilha.auto_filter.ref = f"A1:{get_column_letter(len(CABECALHO))}{ultima}"


def _aba_resumo(planilha, lancamentos: list[Lancamento]) -> None:
    planilha.title = "Resumo"

    receitas = sum(
        (item.valor for item in lancamentos if item.tipo is TipoLancamento.RECEITA),
        Decimal("0.00"),
    )
    despesas = sum(
        (item.valor for item in lancamentos if item.tipo is TipoLancamento.DESPESA),
        Decimal("0.00"),
    )

    _escrever_cabecalho(planilha, ["Indicador", "Valor"])
    for rotulo, valor in (
        ("Receitas", receitas),
        ("Despesas", despesas),
        ("Saldo", receitas - despesas),
        ("Lançamentos", len(lancamentos)),
    ):
        planilha.append([rotulo, valor])

    for linha in range(2, 5):  # as três primeiras são dinheiro; a contagem não
        planilha.cell(row=linha, column=2).number_format = FORMATO_MOEDA

    planilha.append([])
    inicio_categorias = planilha.max_row + 1
    planilha.append(["Categoria", "Tipo", "Total"])
    for coluna in range(1, 4):
        planilha.cell(row=inicio_categorias, column=coluna).font = Font(bold=True)

    por_categoria: dict[tuple[str, str], Decimal] = {}
    for item in lancamentos:
        chave = (item.categoria.nome, item.tipo.rotulo)
        por_categoria[chave] = por_categoria.get(chave, Decimal("0.00")) + item.valor

    for (categoria, tipo), total in sorted(
        por_categoria.items(), key=lambda item: item[1], reverse=True
    ):
        planilha.append([categoria, tipo, total])

    for linha in range(inicio_categorias + 1, planilha.max_row + 1):
        planilha.cell(row=linha, column=3).number_format = FORMATO_MOEDA

    planilha.column_dimensions["A"].width = 24
    planilha.column_dimensions["B"].width = 12
    planilha.column_dimensions["C"].width = 16


def gerar_xlsx(lancamentos: list[Lancamento]) -> bytes:
    livro = Workbook()
    _aba_lancamentos(livro.active, lancamentos)
    _aba_resumo(livro.create_sheet(), lancamentos)

    buffer = io.BytesIO()
    livro.save(buffer)
    return buffer.getvalue()
